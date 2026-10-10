"""
Excel Parser for Indian Railways SLAM Bookings
"""
import os
import re
import openpyxl
from typing import List, Dict, Any, Optional
from config import SECTION_IDS, DEFAULT_OBS_TYPE, DEFAULT_JOB_KIND

KNOWN_SECTIONS = set(SECTION_IDS.keys())

def extract_loco_from_text(text: str) -> Optional[str]:
    """Extract a 5-digit locomotive number (e.g. 30389, 37931)."""
    if not text:
        return None
    m = re.search(r'\b(3\d{4})\b', text) # Electric locos commonly start with 3xxxx
    if m:
        return m.group(1)
    # Generic 5-digit check
    m2 = re.search(r'\b(\d{5})\b', text)
    if m2:
        return m2.group(1)
    return None

import config

def is_all_sec_string(sec_str: str) -> bool:
    """
    Checks if a section string represents 'All Sec' / 'All Sections'.
    Handles variations like 'All Sec', 'ALL SEC', 'All sec.', 'ALL SECTION', 'AllSec', 'ALL-SEC'.
    """
    if not sec_str:
        return False
    s = str(sec_str).strip().upper()
    if s in {"ALL SEC", "ALL SEC.", "ALL SECS", "ALL SECTION", "ALL SECTIONS", "ALL-SEC", "ALLSEC", "ALL"}:
        return True
    if re.search(r'\bALL\s*SEC(?:TION)?S?\.?\b', s):
        return True
    return False

def normalize_sections(sec_str: str, all_sec_custom: Optional[List[str]] = None) -> List[str]:
    """
    Parse strings like 'All Sec', 'E3A & E5A', 'M1 / M2', 'E8 / E4', 'Q&I', 'E4/E5B', 'M2 & E3A'
    into validated SLAM section codes.
    If 'All Sec' is detected, maps to the configured preference (default: all sections except E8, QAI, LAB, PPIO).
    """
    if not sec_str:
        return []

    # 1. Check if the entire string is 'All Sec'
    if is_all_sec_string(sec_str):
        if all_sec_custom:
            return [s for s in all_sec_custom if s in KNOWN_SECTIONS]
        return config.get_all_sec_preferences()

    # Pre-normalize Q&I / Q & I / Q/I to QAI before splitting by '&' or '/'
    processed = re.sub(r'\bQ\s*&\s*I\b', 'QAI', str(sec_str), flags=re.IGNORECASE)
    processed = re.sub(r'\bQ\s*/\s*I\b', 'QAI', processed, flags=re.IGNORECASE)
    
    # Split by common delimiters: &, /, and, comma, plus
    tokens = re.split(r'[/&+,]|\band\b', processed, flags=re.IGNORECASE)
    matched = []
    
    for token in tokens:
        cleaned = token.strip().upper()
        if not cleaned:
            continue

        # Check if this token is 'All Sec'
        if is_all_sec_string(cleaned):
            all_sec_list = all_sec_custom or config.get_all_sec_preferences()
            for s in all_sec_list:
                if s not in matched:
                    matched.append(s)
            continue
            
        # Direct match
        if cleaned in KNOWN_SECTIONS:
            if cleaned not in matched:
                matched.append(cleaned)
        else:
            # Try fuzzy match in known sections
            for ks in sorted(KNOWN_SECTIONS, key=len, reverse=True):
                # E.g. 'E3A' should match before 'E3'
                if re.search(r'\b' + re.escape(ks) + r'\b', cleaned):
                    if ks not in matched:
                        matched.append(ks)
                    break
                    
    # Fallback: if nothing matched but raw text exists
    if not matched:
        raw_upper = processed.strip().upper()
        if is_all_sec_string(raw_upper):
            all_sec_list = all_sec_custom or config.get_all_sec_preferences()
            return list(all_sec_list)
        for ks in sorted(KNOWN_SECTIONS, key=len, reverse=True):
            if ks in raw_upper:
                if ks not in matched:
                    matched.append(ks)
                    
    return matched

def parse_excel_bookings(file_path: str) -> Dict[str, Any]:
    """
    Parses a shed bookings Excel file and returns structured details.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    filename = os.path.basename(file_path)
    loco_from_fn = extract_loco_from_text(filename)
    
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except (PermissionError, IOError, OSError):
        # File might be open with an exclusive lock in Excel; copy to temp file and parse
        import tempfile
        import shutil
        temp_dir = tempfile.gettempdir()
        temp_copy = os.path.join(temp_dir, f"slam_read_{os.getpid()}_{filename}")
        try:
            shutil.copy2(file_path, temp_copy)
            wb = openpyxl.load_workbook(temp_copy, data_only=True)
        finally:
            try:
                if os.path.exists(temp_copy):
                    os.remove(temp_copy)
            except Exception:
                pass
    sheet = wb.active

    loco_number = loco_from_fn
    date_str = ""
    schedule_title = ""
    
    header_row_idx = None
    col_sno = None
    col_desc = None
    col_sec = None

    # Step 1: Scan top 10 rows for Title, Date, and Table Headers
    for r in range(1, min(sheet.max_row + 1, 12)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(sheet.max_column + 1, 10))]
        
        # Check for Date and Title in Row 2 or nearby
        for val in row_vals:
            if val is not None:
                val_s = str(val).strip()
                if "date" in val_s.lower():
                    # e.g. "Date \n 31-08-26" or "Date \n15-09-26"
                    dm = re.search(r'(\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})', val_s)
                    if dm:
                        date_str = dm.group(1)
                if any(kw in val_s.upper() for kw in ["BOOKING", "CHECKING", "QAI", "LBR", "VOLTAGE"]):
                    schedule_title = val_s
                    if not loco_number:
                        loco_number = extract_loco_from_text(val_s)
                        
        # Check if this row looks like the table header row
        r_sno = None
        r_desc = None
        r_sec = None

        for c, val in enumerate(row_vals, start=1):
            if val is None:
                continue
            v_str = str(val).strip()
            v_clean = re.sub(r'[^A-Z0-9]', '', v_str.upper())

            # Match S.No variations: 'SL. NO', 'SL.NO', 'SL NO', 'S.NO', 'SNO', 'SR. NO', 'SERIAL NO', 'NO'
            if v_clean in {'SNO', 'SLNO', 'SRNO', 'SERIALNO', 'NO'} or v_clean.startswith(('SLNO', 'SRNO', 'SNO')):
                r_sno = c
            # Match Section header
            elif 'SECTION' in v_clean:
                r_sec = c
            # Match Description header (skip titles with loco numbers or long working banners)
            elif any(kw in v_clean for kw in ['BOOKING', 'DESCRIPTION', 'DEFECT', 'OBSERVATION', 'REMARK', 'PARTICULAR']):
                if len(v_str) < 30 and not any(k in v_str.upper() for k in ['WORKING', 'PANTO', 'TFP', 'UNDER', 'LOCO']):
                    r_desc = c

        if (r_sno and r_desc) or (r_sno and r_sec) or (r_desc and r_sec):
            header_row_idx = r
            if r_sno: col_sno = r_sno
            if r_desc: col_desc = r_desc
            if r_sec: col_sec = r_sec
            break

    # Fallback 1: Scan for numeric S.No sequence (1, 2, 3...) if col_sno not identified
    if col_sno is None:
        for c in range(1, min(sheet.max_column + 1, 10)):
            for r in range(1, min(sheet.max_row, 12)):
                val1 = sheet.cell(row=r, column=c).value
                val2 = sheet.cell(row=r+1, column=c).value
                try:
                    if int(float(str(val1).strip())) == 1 and int(float(str(val2).strip())) == 2:
                        col_sno = c
                        if header_row_idx is None:
                            header_row_idx = r - 1
                        break
                except Exception:
                    pass
            if col_sno:
                break

    # Fallback 2: Infer missing col_desc or col_sec from first data row
    if header_row_idx and (col_desc is None or col_sec is None):
        data_row = header_row_idx + 1
        for c in range(1, min(sheet.max_column + 1, 10)):
            if c == col_sno:
                continue
            v = str(sheet.cell(row=data_row, column=c).value or '').strip()
            if not v:
                continue
            if any(sec in v.upper() for sec in KNOWN_SECTIONS) and len(v) < 15 and col_sec is None:
                col_sec = c
            elif len(v) > 8 and col_desc is None:
                col_desc = c

    # Ultimate fallback defaults
    if header_row_idx is None:
        header_row_idx = 3
    if col_sno is None:
        col_sno = 1
    if col_desc is None:
        col_desc = 2
    if col_sec is None:
        col_sec = 3

    # Ensure no column index collisions
    if col_sno == col_desc or col_desc == col_sec or col_sno == col_sec:
        if col_sno == 1:
            col_desc = 2
            col_sec = 3
        else:
            col_sno = 2
            col_desc = 3
            col_sec = 4

    # Extract schedule type if detectable in title or filename
    schedule = ""
    KNOWN_SCHEDULES = {
        "IA", "IB", "IC", "GC", "ET", "TI", "TOH", "IOH", "POH", "AOH",
        "MTR", "QAI", "LBR", "PST", "UNSCH", "M", "M1", "M2", "M12", "M24", "M36"
    }

    # Check filename first, e.g. "30389 - GC - 15-09- 26" -> "GC", "37229 - ET - 20-09-26" -> "ET"
    fn_parts = [p.strip() for p in re.split(r'[-_]', filename) if p.strip()]
    for p in fn_parts:
        up = p.upper()
        if up in KNOWN_SCHEDULES:
            schedule = up
            break
            
    if not schedule and schedule_title:
        title_parts = [p.strip() for p in re.split(r'[-_]', schedule_title) if p.strip()]
        for p in title_parts:
            up = p.upper()
            if up in KNOWN_SCHEDULES:
                schedule = up
                break
        if not schedule and len(title_parts) >= 2:
            schedule = title_parts[-1]

    # Smart fallback: short alphabetic token between hyphens
    if not schedule:
        for p in fn_parts:
            up = p.upper()
            if re.match(r'^[A-Z]{2,4}$', up) and not any(kw in up for kw in ["DATE", "BOOK", "XLSX", "SLIP", "TEST", "LOCO", "SHED"]):
                schedule = up
                break

    bookings = []
    current_category = "General"

    # Step 2: Read booking rows
    for r in range(header_row_idx + 1, sheet.max_row + 1):
        sno_val = sheet.cell(row=r, column=col_sno).value
        desc_val = sheet.cell(row=r, column=col_desc).value
        sec_val = sheet.cell(row=r, column=col_sec).value

        if desc_val is None and sec_val is None and sno_val is None:
            continue

        desc_str = str(desc_val).strip() if desc_val is not None else ""
        sec_str = str(sec_val).strip() if sec_val is not None else ""

        # Category headers (e.g., S.No is empty, but description is like "UG / MISC BOOKING")
        if sno_val is None or str(sno_val).strip() == "":
            if desc_str and not sec_str:
                current_category = desc_str
            continue

        # Check if S.No is numeric
        sno_int = None
        try:
            sno_int = int(float(str(sno_val).strip()))
        except (ValueError, TypeError):
            # Not a data row
            continue

        if not desc_str:
            continue

        # Normalize sections
        matched_secs = normalize_sections(sec_str)

        bookings.append({
            "sno": sno_int,
            "description": desc_str,
            "sections": matched_secs,
            "raw_section": sec_str,
            "category": current_category,
            "obs_type": DEFAULT_OBS_TYPE,
            "kind": DEFAULT_JOB_KIND
        })

    return {
        "file_name": filename,
        "loco_number": loco_number or "Unknown",
        "date": date_str,
        "schedule_title": schedule_title,
        "schedule": schedule,
        "bookings_count": len(bookings),
        "bookings": bookings
    }


# ==============================================================================
# LOCO EQUIPMENT EXCEL & GOOGLE SHEET PARSER
# ==============================================================================
import csv
import io
import urllib.request
from datetime import datetime as _dt, date as _date, timedelta as _timedelta
from typing import Tuple


def is_blank_eq_value(value: Any) -> bool:
    """Returns True if value is None, empty, NA, N/A, NIL, -, or NONE."""
    if value is None:
        return True
    v = str(value).strip().upper()
    return v in ("", "NA", "N/A", "NIL", "UNKNOWN", "--SELECT--", "-", "NONE", "NAN", "NULL", "#VALUE!", "#REF!", "#N/A", "#ERROR!", "#NAME?")


def format_excel_scalar(value: Any) -> str:
    """Formats numeric or string Excel values cleanly (e.g. 30541.0 -> '30541', 3699.0 -> '3699')."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "YES" if value else "NO"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    s = str(value).strip()
    if re.match(r"^\d+\.0$", s):
        return s[:-2]
    return s


def format_excel_date(value: Any) -> str:
    """
    Converts Excel datetime objects, Excel serial date floats (e.g. 45727.0 -> '11/03/2025'),
    or date strings into DD/MM/YYYY format expected by SLAM.
    """
    if value is None:
        return ""
    if isinstance(value, _dt):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, _date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        num = float(value)
        if 25000 <= num <= 75000:
            try:
                dt_val = _dt(1899, 12, 30) + _timedelta(days=int(round(num)))
                return dt_val.strftime("%d/%m/%Y")
            except Exception:
                pass
        if num.is_integer():
            return str(int(num))
        return str(value)

    s = str(value).strip()
    if not s or is_blank_eq_value(s) or s.upper() == "NO":
        return "" if is_blank_eq_value(s) else s.upper()

    # Check if string is an Excel serial number like '45727' or '45727.0'
    if re.match(r"^\d{5}(?:\.0+)?$", s):
        try:
            num = float(s)
            if 25000 <= num <= 75000:
                dt_val = _dt(1899, 12, 30) + _timedelta(days=int(round(num)))
                return dt_val.strftime("%d/%m/%Y")
        except Exception:
            pass

    # Convert YYYY-MM-DD HH:MM:SS or YYYY-MM-DD to DD/MM/YYYY
    m_iso = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:\s+.*)?$", s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        return f"{d:02d}/{m:02d}/{y:04d}"

    # Normalize DD-MM-YYYY or DD.MM.YYYY to DD/MM/YYYY
    m_dmy = re.match(r"^(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})$", s)
    if m_dmy:
        d, m, y_raw = int(m_dmy.group(1)), int(m_dmy.group(2)), m_dmy.group(3)
        y = int(y_raw) + 2000 if len(y_raw) == 2 else int(y_raw)
        return f"{d:02d}/{m:02d}/{y:04d}"

    return s


def split_eq_name_and_pos(eq_raw: Any, pos_raw: Any) -> Tuple[str, str, str]:
    """
    Normalizes equipment name and position.
    Examples:
      ('TMB1', None)  -> ('TMB1', 'TMB', '1')
      ('MRB2', 2.0)   -> ('MRB2', 'MRB', '2')
      ('OCB', '1')    -> ('OCB1', 'OCB', '1')
    Returns (eq_display, eq_base_name, position_str).
    """
    eq_s = format_excel_scalar(eq_raw).upper().strip()
    pos_s = format_excel_scalar(pos_raw).strip()
    if is_blank_eq_value(pos_s):
        pos_s = ""

    m = re.match(r"^([A-Z]+)\s*0*(\d+)$", eq_s)
    if m:
        base_name = m.group(1)
        inferred_pos = m.group(2)
        final_pos = pos_s.lstrip("0") if pos_s else inferred_pos
        if not final_pos:
            final_pos = inferred_pos
        return (f"{base_name}{final_pos}", base_name, final_pos)

    if eq_s and pos_s:
        clean_pos = pos_s.lstrip("0") or pos_s
        return (f"{eq_s}{clean_pos}", eq_s, clean_pos)

    return (eq_s, eq_s, pos_s)


MAKE_ALIAS_MAP = {
    "FE": "FLORICON",
    "FLORICON": "FLORICON",
    "BG": "B G IND.",
    "B G": "B G IND.",
    "BG IND": "B G IND.",
    "BG IND.": "B G IND.",
    "B G IND": "B G IND.",
    "B G IND.": "B G IND.",
    "BHL": "BHEL",
    "BHEL": "BHEL",
    "TROLEX": "TROLEX",
    "ALLMETRA": "Allmetra",
    "TOP GRIP": "TOP GRIP",
    "ARIHANT": "ARIHANT",
    "UNKNOWN": "",
    "NIL": "",
}


def normalize_make_alias(make_raw: Any) -> str:
    """Resolves common section shorthand make codes (e.g., FE -> FLORICON, BG -> B G IND., BHL -> BHEL)."""
    s = format_excel_scalar(make_raw).strip()
    if is_blank_eq_value(s) or s.upper() == "NO":
        return ""
    u = re.sub(r"\s+", " ", s.upper()).strip()
    if u in MAKE_ALIAS_MAP:
        return MAKE_ALIAS_MAP[u]
    return s


def normalize_schedule_code(sched_raw: Any) -> str:
    """Extracts a clean SLAM schedule code from raw strings like 'IOH/PER' -> 'IOH', 'POH /lw/per' -> 'POH'."""
    s = format_excel_scalar(sched_raw).strip()
    if is_blank_eq_value(s):
        return ""
    # If it's a date accidentally placed in OH Type column, ignore
    if re.match(r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}$", s):
        return ""
    u = s.upper().strip()
    if "/" in u:
        first_part = u.split("/")[0].strip()
        if first_part in ("IA", "IB", "IC", "IOH", "TOH", "POH", "MTR", "QTR", "TI", "UNSCH", "SPL"):
            return "UnSch" if first_part in ("UNSCH", "SPL") else first_part
    if u in ("COMM", "SPL", "UNSCH", "UN-SCH"):
        return "UnSch"
    return s


def _normalize_eq_dict_row(sno: int, raw_dict: Dict[str, Any], default_eq_name: str = "") -> Optional[Dict[str, Any]]:
    """Normalizes a dictionary of raw header->cell values into a standardized equipment row dict."""
    norm_map = {}
    for k, v in raw_dict.items():
        if k is None:
            continue
        k_clean = re.sub(r"\s+", " ", str(k).strip()).lower()
        norm_map[k_clean] = v

    def _get(*keys):
        for k in keys:
            if k.lower() in norm_map:
                val = norm_map[k.lower()]
                if val is not None and str(val).strip() != "":
                    return val
        return None

    loco_val = _get("loco no", "loco", "loco number", "loco_no")
    loco_no = format_excel_scalar(loco_val)
    # Strip trailing '/1' if loco_no is '37594/1'
    if "/" in loco_no and re.match(r"^\d{5}/\d+$", loco_no):
        loco_no = loco_no.split("/")[0]
    if is_blank_eq_value(loco_no) or not re.match(r"^\d{4,6}$", loco_no):
        return None

    eq_val = _get("eq name", "item", "equipment", "eq.name", "eq_name", "eqpt name", "eqpt. name") or default_eq_name
    pos_val = _get("position", "pos", "eq.position", "tm/s.no", "tm provided no", "tm no")
    eq_raw_disp, eq_name, pos_str = split_eq_name_and_pos(eq_val, pos_val)
    if is_blank_eq_value(eq_name):
        return None

    # Normalize common equipment names from TM details sheets
    eq_u = eq_name.upper().strip()
    if eq_u in ("TEMPERTURE SENSOR", "TEMPERATURE SENSOR", "TEMP SENSOR", "TM TEMP SENSOR"):
        eq_name = "TM TEMP SENSOR"
        eq_raw_disp = "TM TEMP SENSOR"
    elif eq_u in ("SPEED SENSOR",):
        eq_name = "SPEED SENSOR"
        eq_raw_disp = "SPEED SENSOR"

    raw_sched = _get("schedule", "sch", "latest fitment", "latest oh type", "oh type", "oh type/shed")
    schedule = normalize_schedule_code(raw_sched) or "IA"

    prov_raw = _get("provision date", "prov date", "provision_date", "provided date", "latest fitment date", "l/date", "oh date", "oh  date", "update date till", "update till", "date")
    prov_date = format_excel_date(prov_raw)
    if not re.match(r"^\d{2}/\d{2}/\d{4}$", prov_date):
        # Fallback if latest fitment date was in OH Type/SHED column
        alt_prov = format_excel_date(_get("oh type/shed", "oh date", "commdt", "comm. dt"))
        prov_date = alt_prov if re.match(r"^\d{2}/\d{2}/\d{4}$", alt_prov) else "03/09/2026"

    model = format_excel_scalar(_get("model", "eq.model", "eq model"))
    if is_blank_eq_value(model):
        model = "NA"

    eq_num = format_excel_scalar(_get("eq. number (m)", "eq. number", "eq.number", "sr. no.", "sr no", "serial no", "slno", "sl.no."))
    if is_blank_eq_value(eq_num):
        eq_num = f"{loco_no}-TM{pos_str}" if pos_str else f"{loco_no}-1"

    loc_eq_num = format_excel_scalar(_get("location eq. number", "eq. number (l)", "location eq.number", "loc eq. no"))
    if is_blank_eq_value(loc_eq_num):
        loc_eq_num = eq_num or (f"{eq_raw_disp}" if eq_raw_disp in ("TMB1", "TMB2", "MRB1", "MRB2", "OCB1", "OCB2") else (pos_str or "1"))

    doc = format_excel_date(_get("doc", "date of commission", "commdt", "comm. dt"))
    if not re.match(r"^\d{2}/\d{2}/\d{4}$", doc):
        doc = prov_date or "09/09/2026"
    dom = format_excel_date(_get("dom", "date of manufacture"))
    if not re.match(r"^\d{2}/\d{2}/\d{4}$", dom):
        dom = "10/09/2026"
    rec_date = format_excel_date(_get("date of receive", "date of recieve", "receive date"))
    if not re.match(r"^\d{2}/\d{2}/\d{4}$", rec_date):
        rec_date = prov_date or "03/09/2026"

    make = normalize_make_alias(_get("make", "eq.make", "eq make")) or "FLORICON"

    warranty = format_excel_scalar(_get("warranty")).upper()
    if warranty not in ("YES", "NO"):
        warranty = "YES"

    warranty_upto = format_excel_date(_get("warranty upto", "warranty upto(if avbl)"))
    if warranty_upto.upper() == "NO" or not re.match(r"^\d{2}/\d{2}/\d{4}$", warranty_upto):
        warranty_upto = "20/09/2028" if warranty == "YES" else "NO"

    amc = format_excel_scalar(_get("amc")).upper()
    if amc not in ("YES", "NO"):
        amc = "NO"

    amc_upto = format_excel_date(_get("amc upto", "amc upto (if avbl)"))
    if amc_upto.upper() == "NO" or not re.match(r"^\d{2}/\d{2}/\d{4}$", amc_upto):
        amc_upto = "NO"

    remarks = format_excel_scalar(_get("remarks", "remark", "eqpt remarks"))
    if is_blank_eq_value(remarks):
        remarks = "OK"

    return {
        "sno": sno,
        "loco_no": loco_no,
        "eq_raw": eq_raw_disp,
        "eq_name": eq_name,
        "position": pos_str,
        "schedule": schedule,
        "provision_date": prov_date,
        "model": model,
        "eq_number": eq_num,
        "loc_eq_number": loc_eq_num,
        "doc": doc,
        "dom": dom,
        "receive_date": rec_date,
        "make": make,
        "warranty": warranty,
        "warranty_upto": warranty_upto,
        "amc": amc,
        "amc_upto": amc_upto,
        "remarks": remarks,
    }


def _infer_matrix_eq_name(sheet_title: str, ws) -> str:
    """Infers the SLAM Equipment Name for a wide TM1..TM6 matrix sheet."""
    st = sheet_title.strip().upper()
    if "TEMP" in st:
        return "TM TEMP SENSOR"
    if "SPEED" in st:
        return "SPEED SENSOR"
    if "PINION" in st:
        return "PINION"
    if "ROTOR" in st:
        return "ROTOR"
    if "BRG" in st or "BEAR" in st:
        return "BEARING"
    # Check top 4 rows of sheet for a title banner
    for r in range(1, min(ws.max_row + 1, 5)):
        for c in range(1, min(ws.max_column + 1, 10)):
            val = ws.cell(r, c).value
            if val:
                vu = str(val).strip().upper()
                if "TEMP" in vu and "SENSOR" in vu:
                    return "TM TEMP SENSOR"
                if "SPEED" in vu and "SENSOR" in vu:
                    return "SPEED SENSOR"
    return st or "TM EQUIPMENT"


def _parse_wide_tm_matrix_sheet(ws, sheet_title: str, header_row: int) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Parses wide matrix sheets (like 'Temperture sensor', 'speed sensor', 'PINION', 'BRG MAKE', 'ROTOR')
    where each row has a Loco No and columns TM1..TM6 representing Positions 1..6.
    """
    eq_name = _infer_matrix_eq_name(sheet_title, ws)
    max_col = min(ws.max_column, 30)
    row_vals = [ws.cell(header_row, c).value for c in range(1, max_col + 1)]
    next_row_vals = [ws.cell(header_row + 1, c).value for c in range(1, max_col + 1)] if header_row < ws.max_row else [None] * max_col

    loco_col = None
    date_col = None
    tm_cols: Dict[str, Dict[str, int]] = {}  # pos -> {"make": col_idx, "model": col_idx}

    for c_idx in range(1, max_col + 1):
        v_str = str(row_vals[c_idx - 1] or "").strip().upper()
        v_next = str(next_row_vals[c_idx - 1] or "").strip().upper()
        if "LOCO" in v_str and loco_col is None:
            loco_col = c_idx
        elif ("UPDATE" in v_str or "DATE" in v_str) and date_col is None:
            date_col = c_idx
        elif ("UPDATE" in v_next or "DATE" in v_next) and date_col is None:
            date_col = c_idx

        m_tm = re.match(r"^TM\s*([1-6])$", v_str)
        if m_tm:
            pos = m_tm.group(1)
            tm_cols[pos] = {"make": c_idx}
            # Check if next row has Make / Model sub-headers (like ROTOR sheet)
            if v_next == "MAKE" and c_idx < max_col:
                v_next_right = str(next_row_vals[c_idx] or "").strip().upper()
                if v_next_right == "MODEL":
                    tm_cols[pos]["model"] = c_idx + 1

    # In case Loco No was 1 row above TM1..TM6 (like 'pinion gap' sheet)
    if loco_col is None and header_row > 1:
        for c_idx in range(1, max_col + 1):
            v_prev = str(ws.cell(header_row - 1, c_idx).value or "").strip().upper()
            if "LOCO" in v_prev:
                loco_col = c_idx
                break

    if loco_col is None or not tm_cols:
        return [], []

    rows: List[Dict[str, Any]] = []
    locos: List[str] = []
    loco_set = set()

    for r in range(header_row + 1, ws.max_row + 1):
        raw_loco = ws.cell(r, loco_col).value
        loco_no = format_excel_scalar(raw_loco)
        if not re.match(r"^\d{4,6}$", loco_no):
            continue

        prov_date = ""
        if date_col:
            cand_dt = format_excel_date(ws.cell(r, date_col).value)
            if re.match(r"^\d{2}/\d{2}/\d{4}$", cand_dt):
                prov_date = cand_dt

        for pos in sorted(tm_cols.keys(), key=lambda x: int(x)):
            c_info = tm_cols[pos]
            raw_make = ws.cell(r, c_info["make"]).value
            make_val = normalize_make_alias(raw_make)
            model_val = ""
            if "model" in c_info:
                raw_mod = format_excel_scalar(ws.cell(r, c_info["model"]).value)
                if not is_blank_eq_value(raw_mod):
                    model_val = raw_mod

            # Skip completely empty/unfitted positions in matrix
            if not make_val and not model_val:
                continue

            eff_prov = prov_date or "03/09/2026"
            sno = len(rows) + 1
            rows.append({
                "sno": sno,
                "loco_no": loco_no,
                "eq_raw": eq_name,
                "eq_name": eq_name,
                "position": pos,
                "schedule": "IA",
                "provision_date": eff_prov,
                "model": model_val or "NA",
                "eq_number": f"{loco_no}-TM{pos}",
                "loc_eq_number": f"{loco_no}-TM{pos}",
                "doc": "09/09/2026",
                "dom": "10/09/2026",
                "receive_date": eff_prov,
                "make": make_val,
                "warranty": "YES",
                "warranty_upto": "20/09/2028",
                "amc": "NO",
                "amc_upto": "NO",
                "remarks": "OK",
            })
            if loco_no not in loco_set:
                loco_set.add(loco_no)
                locos.append(loco_no)

    return rows, locos


def parse_equipment_excel(file_path: str, sheet_name: Optional[str] = None) -> Dict[str, Any]:
    """
    Parses a Loco Equipment Excel file — supporting BOTH:
      1. Standard formatted templates ('equipment fill templete.xlsx')
      2. Raw multi-sheet section workbooks ('TM in & out details.xlsx' with wide TM1..TM6 matrices
         like 'Temperture sensor', 'speed sensor', 'PINION', 'ROTOR', 'BRG MAKE', or tabular logs
         like 'TM details from 2301' and 'TM details from 82024').
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Equipment Excel file not found: {file_path}")

    filename = os.path.basename(file_path)
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except (PermissionError, IOError, OSError):
        import tempfile
        import shutil
        temp_copy = os.path.join(tempfile.gettempdir(), f"slam_eq_{os.getpid()}_{filename}")
        try:
            shutil.copy2(file_path, temp_copy)
            wb = openpyxl.load_workbook(temp_copy, data_only=True)
        finally:
            try:
                if os.path.exists(temp_copy):
                    os.remove(temp_copy)
            except Exception:
                pass

    sheet_names = wb.sheetnames
    matched_sheet = None
    if sheet_name:
        if sheet_name in sheet_names:
            matched_sheet = sheet_name
        else:
            sn_clean = sheet_name.strip().lower()
            for s in sheet_names:
                if s.strip().lower() == sn_clean:
                    matched_sheet = s
                    break
    if matched_sheet:
        sheet_name = matched_sheet
        ws = wb[sheet_name]
    else:
        # If opening a multi-sheet TM file without explicit sheet_name, prefer 'Temperture sensor' if present
        preferred = None
        for s in sheet_names:
            if "temperture sensor" in s.lower() or "temperature sensor" in s.lower():
                preferred = s
                break
        sheet_name = preferred or sheet_names[0]
        ws = wb[sheet_name]

    # Scan top 8 rows to detect whether this sheet is a Wide TM1..TM6 Matrix or a Tabular Sheet
    matrix_header_row = None
    tabular_header_row = None
    headers: List[str] = []

    for r in range(1, min(ws.max_row + 1, 9)):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 28))]
        joined_upper = " ".join(str(v).strip().upper() for v in vals if v is not None)
        if "TM1" in joined_upper and "TM2" in joined_upper and "TM3" in joined_upper:
            matrix_header_row = r
            break
        if "LOCO" in joined_upper and any(k in joined_upper for k in ("EQ", "ITEM", "MAKE", "SLNO", "TM PROVIDED")):
            tabular_header_row = r
            headers = [str(v).strip() if v is not None else f"col_{i}" for i, v in enumerate(vals, start=1)]
            break

    if matrix_header_row is not None:
        rows, locos = _parse_wide_tm_matrix_sheet(ws, sheet_name, matrix_header_row)
        return {
            "source": file_path,
            "file_name": filename,
            "sheet_name": sheet_name,
            "sheet_names": sheet_names,
            "format_type": "raw_matrix",
            "locos": locos,
            "locos_count": len(locos),
            "rows_count": len(rows),
            "rows": rows,
        }

    header_row = tabular_header_row or 1
    if not headers:
        headers = [
            str(ws.cell(header_row, c).value).strip() if ws.cell(header_row, c).value is not None else f"col_{c}"
            for c in range(1, min(ws.max_column + 1, 25))
        ]

    default_eq = "NBC BEARING" if "nbc" in sheet_name.lower() else ""
    rows = []
    locos = []
    loco_set = set()

    for r in range(header_row + 1, ws.max_row + 1):
        raw_dict = {}
        for c_idx, h_name in enumerate(headers, start=1):
            try:
                raw_dict[h_name] = ws.cell(r, c_idx).value
            except Exception:
                raw_dict[h_name] = None
        item = _normalize_eq_dict_row(len(rows) + 1, raw_dict, default_eq_name=default_eq)
        if item:
            if default_eq == "NBC BEARING" and not item.get("make"):
                item["make"] = "NBC"
            rows.append(item)
            lno = item["loco_no"]
            if lno not in loco_set:
                loco_set.add(lno)
                locos.append(lno)

    # Guard against Excel drag-fill on the 'position' column (e.g. 1..4 on first loco, 5..8 on second loco, etc.)
    from collections import defaultdict
    grouped_by_loco_eq = defaultdict(list)
    for r_item in rows:
        grouped_by_loco_eq[(r_item["loco_no"], r_item["eq_name"])].append(r_item)
    for (_, _), grp in grouped_by_loco_eq.items():
        if len(grp) >= 2:
            pos_ints = [int(x["position"]) for x in grp if str(x.get("position", "")).isdigit()]
            if pos_ints and (max(pos_ints) > 6 or (min(pos_ints) > 1 and len(pos_ints) == len(grp))):
                for idx_in_grp, r_item in enumerate(grp, start=1):
                    r_item["position"] = str(idx_in_grp)

    return {
        "source": file_path,
        "file_name": filename,
        "sheet_name": sheet_name,
        "sheet_names": sheet_names,
        "format_type": "tabular",
        "locos": locos,
        "locos_count": len(locos),
        "rows_count": len(rows),
        "rows": rows,
    }


def export_equipment_to_excel(rows: List[Dict[str, Any]], output_path: str, sheet_title: str = "Equipment Fill") -> str:
    """
    Exports equipment rows to a clean, formatted Excel (.xlsx) file matching 'equipment fill templete.xlsx'.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers = [
        "Loco No",
        "Eq Name",
        "position",
        "Schedule",
        "Provision Date",
        "Model",
        "Eq. Number (M)",
        "Location Eq. Number",
        "DOC",
        "DOM",
        "Date of Receive",
        "Make",
        "Warranty",
        "Warranty Upto",
        "AMC",
        "AMC Upto",
        "Remarks",
    ]
    ws.append(headers)

    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([
            r.get("loco_no", ""),
            r.get("eq_raw") or r.get("eq_name", ""),
            r.get("position", ""),
            r.get("schedule", ""),
            r.get("provision_date", ""),
            r.get("model", ""),
            r.get("eq_number", ""),
            r.get("loc_eq_number", ""),
            r.get("doc", ""),
            r.get("dom", ""),
            r.get("receive_date", ""),
            r.get("make", ""),
            r.get("warranty", "NO"),
            r.get("warranty_upto", ""),
            r.get("amc", "NO"),
            r.get("amc_upto", ""),
            r.get("remarks", ""),
        ])

    wb.save(output_path)
    return output_path


def extract_sheet_id_and_gid(url_or_id: str, default_gid: str = "") -> Tuple[str, str]:
    """Extracts Google Sheet ID and GID from either a full Google Docs URL or raw Sheet ID."""
    s = (url_or_id or "").strip()
    gid = default_gid.strip()
    m_id = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", s)
    sheet_id = m_id.group(1) if m_id else s
    m_gid = re.search(r"[#&?]gid=(\d+)", s)
    if m_gid:
        gid = m_gid.group(1)
    return sheet_id, gid


def parse_equipment_google_sheet(url_or_id: str, gid: str = "") -> Dict[str, Any]:
    """
    Downloads and parses a Google Sheet CSV export into the structured equipment format.
    """
    sheet_id, resolved_gid = extract_sheet_id_and_gid(url_or_id, gid)
    if not sheet_id:
        raise ValueError("Please provide a valid Google Sheet URL or Sheet ID.")

    if resolved_gid:
        export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={resolved_gid}"
    else:
        export_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"

    req = urllib.request.Request(
        export_url,
        headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    )
    with urllib.request.urlopen(req, timeout=20) as resp:
        csv_text = resp.read().decode("utf-8-sig", errors="replace")

    reader = csv.DictReader(io.StringIO(csv_text))
    rows: List[Dict[str, Any]] = []
    locos: List[str] = []
    loco_set = set()

    for raw_row in reader:
        item = _normalize_eq_dict_row(len(rows) + 1, raw_row)
        if item:
            rows.append(item)
            lno = item["loco_no"]
            if lno not in loco_set:
                loco_set.add(lno)
                locos.append(lno)

    return {
        "source": export_url,
        "file_name": f"GoogleSheet_{sheet_id[:8]}",
        "sheet_name": f"GID:{resolved_gid or 'default'}",
        "sheet_names": [f"GID:{resolved_gid or 'default'}"],
        "format_type": "google_sheet",
        "locos": locos,
        "locos_count": len(locos),
        "rows_count": len(rows),
        "rows": rows,
    }


def _equipment_names_compatible(portal_eq_name: str, row_eq_name: str, row_eq_raw: str) -> bool:
    """
    Checks whether a SLAM portal equipment name (e.g., 'TM TEMP SENSOR', 'Filter for TMB (Mechanical)')
    matches an Excel equipment row name ('TM TEMP SENSOR', 'TEMP SENSOR', 'TMB', 'TMB1', 'MRB2', etc.).
    """
    p_clean = re.sub(r"\s+", " ", (portal_eq_name or "").upper()).strip()
    r_name = re.sub(r"\s+", " ", (row_eq_name or "").upper()).strip()
    r_raw = re.sub(r"\s+", " ", (row_eq_raw or "").upper()).strip()

    # Strip trailing digits from r_raw/r_name for base comparison (e.g. TMB1 -> TMB)
    r_base = re.sub(r"\s*0*\d+$", "", r_name).strip() or r_name

    if p_clean in (r_name, r_raw, r_base):
        return True

    # TM Temp Sensor variations
    temp_aliases = {"TM TEMP SENSOR", "TEMP SENSOR", "TEMPERTURE SENSOR", "TEMPERATURE SENSOR", "TM TEMPERATURE SENSOR"}
    if p_clean in temp_aliases and (r_name in temp_aliases or r_raw in temp_aliases or r_base in temp_aliases):
        return True

    # Filter for MRB / OCB / TMB variations
    for code in ("TMB", "MRB", "OCB"):
        if code in p_clean and (r_base == code or code in r_name or code in r_raw):
            return True

    # Roof equipment variations
    roof_aliases = {"ROOF", "ROOF EQUIPMENT", "ROOF HATCH", "ROOF ASSEMBLY", "ROOF BUSHING", "FILTER FOR ROOF"}
    if p_clean in roof_aliases and (r_name in roof_aliases or r_raw in roof_aliases or r_base in roof_aliases):
        return True
    if "ROOF" in p_clean and ("ROOF" in r_name or "ROOF" in r_raw or "ROOF" in r_base):
        return True

    # Substring containment (e.g. 'SPEED SENSOR' in 'TM SPEED SENSOR')
    if r_base and len(r_base) >= 3 and (r_base in p_clean or p_clean in r_base):
        return True

    return False


def find_matching_equipment_row(
    rows: List[Dict[str, Any]],
    loco_no: str,
    eq_name: str,
    position: str
) -> Optional[Dict[str, Any]]:
    """
    Matches (loco_no, eq_name, position) against normalized equipment rows,
    handling both separate (eq_name='TM TEMP SENSOR', position='2') and combined ('TMB1') formats.
    """
    target_loco = str(loco_no).strip()
    target_eq = str(eq_name).upper().strip()
    target_pos = str(position).strip().lstrip("0") or str(position).strip()

    compatible_for_loco = []
    for r in rows:
        if str(r.get("loco_no", "")).strip() != target_loco:
            continue
        r_eq = str(r.get("eq_name", "")).upper().strip()
        r_raw = str(r.get("eq_raw", "")).upper().strip()
        r_pos = str(r.get("position", "")).strip().lstrip("0") or str(r.get("position", "")).strip()

        if _equipment_names_compatible(target_eq, r_eq, r_raw):
            compatible_for_loco.append(r)
            if r_pos == target_pos or not target_pos or not r_pos:
                return r

        for cand in (r_raw, r_eq):
            m = re.match(r"^([A-Z\s]+?)\s*0*(\d+)$", cand)
            if m and _equipment_names_compatible(target_eq, m.group(1), cand) and m.group(2) == target_pos:
                return r

    # Fallback: if positions in the sheet were drag-filled (e.g. 5, 6, 7, 8 instead of 1, 2, 3, 4),
    # match the N-th compatible row for this locomotive
    if compatible_for_loco and target_pos.isdigit():
        idx = int(target_pos) - 1
        if 0 <= idx < len(compatible_for_loco):
            return compatible_for_loco[idx]

    return None


